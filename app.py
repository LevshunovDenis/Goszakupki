from datetime import date
from io import BytesIO
from pathlib import Path

from flask import Flask, render_template, request, jsonify, send_file
from bs4 import BeautifulSoup
from docx import Document
from num2words import num2words
import truststore
truststore.inject_into_ssl()
import requests
import re
import urllib.parse

app = Flask(__name__)


USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/127.0.0.0 Safari/537.36"
WORD_TEMPLATE_PATH = Path(__file__).resolve().with_name("word_template.docx")
PURCHASE_TYPES = {
    "etrade": "открытом конкурсе",
    "request": "запросе ценовых предложений",
    "single-source": "закупке из одного источника",
    "marketing": "закупке из одного источника",
    "auction": "электронном аукционе",
    "other": "ином виде процедуры закупки",
}

PURCHASE_TYPE_ALIASES = {
    "etrade": ("etrade", "open-competition", "open", "competition", "конкурс"),
    "request": ("request", "price-request", "proposal", "prices", "ценовых"),
    "single-source": ("single-source", "single_source", "one-source", "one_source", "marketing", "market", "source"),
    "auction": ("auction", "electronic-auction", "auctions", "аукцион"),
}

FIELD_ALIASES = {
    "organizer": (
        "полное наименование организатора, место нахождения организации, унп",
        "наименование закупающей организации",
        "наименование организации",
        "наименование заказчика(-ов) (фио - для ип)",
        "наименование заказчика",
        "заказчик",
    ),
    "subject": (
        "название запроса ценовых предложений",
        "название процедуры закупки",
        "название процедуры закупки из одного источника на этп",
        "название открытого конкурса/конкурса",
        "название закупки",
        "предмет закупки",
        "предмет",
    ),
    "delivery": (
        "срок поставки",
        "срок выполнения",
        "срок поставки товара",
    ),
}


def clean_text(value):
    if value is None:
        return None
    if not isinstance(value, str):
        value = str(value)
    text = re.sub(r"\s+", " ", value).strip()
    text = re.sub(r"[\x00-\x1f\x7f]", " ", text)
    return text[:500] if text else ""


def safe_string(value, max_len=300):
    text = clean_text(value)
    if text is None:
        return ""
    return text[:max_len]


def normalize_key(value):
    if value is None:
        return ""
    return re.sub(r"[^a-zа-яё0-9]+", " ", str(value).lower()).strip()


def pick_first_value(table_data, aliases):
    normalized_table = {normalize_key(key): value for key, value in table_data.items()}
    for alias in aliases:
        candidate = normalized_table.get(normalize_key(alias))
        if candidate:
            return candidate

    normalized_aliases = [normalize_key(alias) for alias in aliases]
    for key, value in normalized_table.items():
        if any(alias in key for alias in normalized_aliases):
            return value
    return None


def currency_word_form(value, forms):
    last_two_digits = value % 100
    if 11 <= last_two_digits <= 14:
        return forms[2]
    last_digit = value % 10
    if last_digit == 1:
        return forms[0]
    if 2 <= last_digit <= 4:
        return forms[1]
    return forms[2]


def amount_to_byn_words(raw_amount):
    amount = raw_amount.replace("\u00a0", " ").replace("\u202f", " ").strip()
    amount = re.sub(
        r"\s*(?:BYN|(?:бел\.?\s*)?руб(?:ль|ля|лей)?\.?)$",
        "",
        amount,
        flags=re.IGNORECASE,
    )
    amount = re.sub(r"\s+", "", amount)
    if not re.fullmatch(r"\d+(?:[.,]\d{1,2})?", amount):
        raise ValueError("Введите сумму числом, например 100 или 700 000,00 BYN.")

    whole_text, separator, kopecks_text = amount.replace(",", ".").partition(".")
    whole = int(whole_text)
    kopecks = int(kopecks_text.ljust(2, "0")) if separator else 0
    rubles_words = num2words(whole, lang="ru")
    ruble_form = currency_word_form(whole, ("рубль", "рубля", "рублей"))
    kopeck_form = currency_word_form(kopecks, ("копейка", "копейки", "копеек"))
    return f"{rubles_words} {ruble_form}, {kopecks:02d} {kopeck_form}"


def normalize_url(raw_url):
    value = (raw_url or "").strip()
    if not value:
        raise ValueError("Введите ссылку или номер закупки")

    if value.isdigit():
        return f"https://goszakupki.by/request/view/{value}"

    if "http" not in value:
        value = "https://" + value

    parsed = urllib.parse.urlparse(value)
    if "goszakupki.by" not in parsed.netloc.lower():
        raise ValueError("Ссылка должна вести на goszakupki.by")

    if parsed.scheme not in {"http", "https"}:
        raise ValueError("Неверная ссылка. Укажите корректный URL на goszakupki.by")

    if re.fullmatch(r"/(?:[A-Za-z0-9_-]+/)?view/\d+/?", parsed.path) or re.fullmatch(r"/(?:[A-Za-z0-9_-]+)?/\d+/?", parsed.path):
        return f"{parsed.scheme}://{parsed.netloc}{parsed.path}"

    raise ValueError("Неверная ссылка. Укажите карточку закупки на goszakupki.by")


def purchase_type_from_url(url):
    parsed = urllib.parse.urlparse(url or "")
    route_type = parsed.path.strip("/").split("/", 1)[0]
    if not route_type:
        return ""

    normalized_route = normalize_key(route_type)
    for purchase_key, aliases in PURCHASE_TYPE_ALIASES.items():
        for alias in aliases:
            if normalize_key(alias) == normalized_route:
                return PURCHASE_TYPES.get(purchase_key, "")

    return PURCHASE_TYPES.get("other", "")


def extract_data_from_html(html):
    soup = BeautifulSoup(html, "html.parser")
    tender_heading = soup.select_one(".breadcrumb .active, .page-header h1")
    tender_number = None
    if tender_heading:
        number_match = re.search(r"\bauc\d+\b", tender_heading.get_text(" ", strip=True), re.IGNORECASE)
        if number_match:
            tender_number = number_match.group(0)

    table_data = {}
    for tr in soup.select("table.table-striped tr"):
        th = tr.find("th")
        td = tr.find("td")
        if th and td:
            key = clean_text(th.get_text(" ", strip=True))
            value = clean_text(td.get_text(" ", strip=True))
            if key:
                table_data[key] = value

    organizer = pick_first_value(table_data, FIELD_ALIASES["organizer"])
    if organizer:
        organizer = re.split(r"\s*Республика Беларусь\b", organizer, maxsplit=1)[0].strip()

    customer = organizer or pick_first_value(table_data, FIELD_ALIASES["organizer"])
    customer = safe_string(customer)
    item = None
    quantity = None
    amount = None
    delivery = None

    lot_desc_el = soup.select_one(".lot-description")
    if lot_desc_el:
        item = safe_string(lot_desc_el.get_text(" ", strip=True))

    count_price_el = soup.select_one(".lot-count-price")
    if count_price_el:
        count_price = safe_string(count_price_el.get_text(" ", strip=True))
        quantity_text, separator, amount_text = count_price.rpartition(", ")
        if separator and re.match(r"\d", amount_text):
            quantity = quantity_text.strip()
            amount = amount_text.strip()
        else:
            quantity = count_price

    for li in soup.select(".lot-inf .list-group-item"):
        text = clean_text(li.get_text(" ", strip=True))
        if any(keyword in text.lower() for keyword in ("срок поставки", "срок выполнения")):
            match = re.search(r"(?:Срок поставки|Срок выполнения)\s*:?\s*(.+)", text)
            if match:
                delivery = match.group(1).strip()
                break

    if not delivery:
        delivery = pick_first_value(table_data, FIELD_ALIASES["delivery"])

    if not item:
        item = pick_first_value(table_data, FIELD_ALIASES["subject"])

    return {
        "tender_number": tender_number,
        "customer": customer or None,
        "subject": safe_string(item) or None,
        "quantity": safe_string(quantity) or None,
        "amount": safe_string(amount) or None,
        "delivery": safe_string(delivery) or None,
    }


def replace_placeholders(container, replacements):
    for paragraph in container.paragraphs:
        for run in paragraph.runs:
            for placeholder, value in replacements.items():
                if placeholder in run.text:
                    run.text = run.text.replace(placeholder, value)

    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                replace_placeholders(cell, replacements)


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/health")
def health_check():
    return jsonify({"status": "ok", "app": "goszakupki_parser"}), 200


@app.route("/api/extract", methods=["POST"])
def extract():
    payload = request.get_json(silent=True) or {}
    raw_url = (payload.get("url") or "").strip()

    try:
        url = normalize_url(raw_url)
    except ValueError as exc:
        return jsonify({"success": False, "error": str(exc)}), 400

    try:
        response = requests.get(
            url,
            timeout=30,
            headers={"User-Agent": USER_AGENT, "Accept-Language": "ru-RU,ru;q=0.9,en;q=0.8"},
            allow_redirects=True,
        )
        response.raise_for_status()
    except requests.RequestException as exc:
        return jsonify({"success": False, "error": f"Не удалось открыть страницу: {exc}"}), 502

    if "text/html" not in response.headers.get("Content-Type", "") and "application/xhtml" not in response.headers.get("Content-Type", ""):
        return jsonify({"success": False, "error": "Получен неожиданный формат ответа сервера."}), 502

    data = extract_data_from_html(response.text)
    data["purchase_type"] = purchase_type_from_url(url)

    if not data.get("customer") and not data.get("subject"):
        return jsonify({"success": False, "error": "На странице не удалось найти нужные данные. Проверьте ссылку."}), 404

    return jsonify({"success": True, "data": data})


@app.route("/api/generate", methods=["POST"])
def generate_document():
    payload = request.get_json(silent=True) or {}
    data = payload.get("data")
    manual_amount = (payload.get("manual_amount") or "").strip()

    if not isinstance(data, dict) or not data:
        return jsonify({"success": False, "error": "Сначала получите данные закупки."}), 400
    if not manual_amount:
        return jsonify({"success": False, "error": "Введите сумму для документа."}), 400
    if len(manual_amount) > 100:
        return jsonify({"success": False, "error": "Слишком длинная сумма для документа."}), 400
    try:
        manual_amount_words = amount_to_byn_words(manual_amount)
    except ValueError as exc:
        return jsonify({"success": False, "error": str(exc)}), 400
    if not WORD_TEMPLATE_PATH.is_file():
        return jsonify({"success": False, "error": "Поместите Word-шаблон word_template.docx рядом с app.py."}), 400

    try:
        document = Document(WORD_TEMPLATE_PATH)
    except Exception:
        return jsonify({"success": False, "error": "Не удалось открыть word_template.docx. Проверьте файл шаблона."}), 400

    today = date.today()
    replacements = {
        "{{DATE_TODAY}}": today.strftime("%d.%m.%Y"),
        "{{TENDER_NUMBER}}": str(data.get("tender_number") or ""),
        "{{PURCHASE_TYPE}}": str(data.get("purchase_type") or ""),
        "{{CUSTOMER}}": str(data.get("customer") or ""),
        "{{SUBJECT}}": str(data.get("subject") or ""),
        "{{QUANTITY}}": str(data.get("quantity") or ""),
        "{{SOURCE_AMOUNT}}": str(data.get("amount") or ""),
        "{{DELIVERY}}": str(data.get("delivery") or ""),
        "{{MY_AMOUNT}}": manual_amount,
        "{{MY_AMOUNT_WORDS}}": manual_amount_words,
    }

    replace_placeholders(document, replacements)
    for section in document.sections:
        replace_placeholders(section.header, replacements)
        replace_placeholders(section.footer, replacements)

    output = BytesIO()
    document.save(output)
    output.seek(0)
    tender_number = re.sub(r"[^A-Za-z0-9_-]", "", str(data.get("tender_number") or ""))
    download_name = f"{tender_number}.docx" if tender_number else f"Закупка_{today:%Y%m%d}.docx"

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        as_attachment=True,
        download_name=download_name,
    )


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5001, debug=False)
